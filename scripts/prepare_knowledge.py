# /// script
# requires-python = ">=3.12"
# dependencies = ["python-pptx", "python-docx", "pymupdf"]
# ///
"""Builds the overlay's knowledge base from the seminar materials.

Usage:
    uv run scripts/prepare_knowledge.py SLIDES.pptx QA.docx THESIS.pdf [OUT_DIR]

Writes OUT_DIR (default: ./knowledge):
    qa.json     prepared question/answer pairs (the Q&A document plus the
                anticipated questions in the slide notes)
    slides.md   slide text with speaker notes
    thesis.md   dissertation body with [p. N] page markers, without the
                reference list and the source-code appendix
"""

import json
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation

# Markers that appear in the Ukrainian source documents.
NOTES_QUESTION = "Можуть запитати"
ANSWER = "Відповідь"
NOTES_SECTION_ICONS = "❓⏱🎯💡"
THESIS_SECTIONS = [  # (first heading, heading that ends the section)
    ("АНОТАЦІЯ", "ABSTRACT"),
    ("СПИСОК ПУБЛІКАЦІЙ ЗДОБУВАЧА", "ЗМІСТ"),
    ("ВСТУП", "СПИСОК ВИКОРИСТАНИХ ДЖЕРЕЛ"),
    ("ДОДАТОК Б", "ДОДАТОК В"),
]

QUESTION_IN_NOTES = re.compile(
    rf"❓\s*{NOTES_QUESTION}:\s*«(.+?)»\s*\n\s*{ANSWER}:\s*(.+?)(?=\n\s*[{NOTES_SECTION_ICONS}]|\Z)", re.S
)
QUESTION_IN_DOC = re.compile(r"^\s*\d+\.\s*(.+?)(?:\s*\(([^()]+)\))?\s*$")


def extract_slides(path):
    pairs, parts = [], []
    for number, slide in enumerate(Presentation(path).slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                texts.extend(" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        body = "\n".join(t for t in texts if t != "‹#›")
        parts.append(f"## Slide {number}\n{body}\n\n### Speaker notes\n{notes}\n")
        for question, answer in QUESTION_IN_NOTES.findall(notes):
            pairs.append({"q": question.strip(), "a": answer.strip(), "source": f"slide {number} notes"})
    return "\n".join(parts), pairs


def extract_qa_doc(path):
    paragraphs = [p.text.strip() for p in Document(path).paragraphs if p.text.strip()]
    pairs, question, asker = [], None, None
    for text in paragraphs:
        if text.startswith(ANSWER):
            if question:
                answer = re.sub(rf"^{ANSWER}\.?\s*", "", text)
                pairs.append({"q": question, "a": answer, "who": asker, "source": "Q&A document"})
            question = None
        elif m := QUESTION_IN_DOC.match(text):
            question, asker = m.group(1).strip(), m.group(2)
    return pairs


def is_heading(line, title):
    # Table-of-contents entries carry dot leaders and page numbers; real headings don't.
    return line.strip().startswith(title) and "..." not in line


def section(lines, start_title, end_title):
    begin = next((i for i, line in enumerate(lines) if is_heading(line, start_title)), None)
    if begin is None:
        return []
    end = next((i for i in range(begin + 1, len(lines)) if is_heading(lines[i], end_title)), len(lines))
    page = next((lines[i] for i in range(begin, -1, -1) if lines[i].startswith("[p. ")), None)
    return ([page] if page else []) + lines[begin:end]


def extract_thesis(path):
    lines = []
    for number, page in enumerate(pymupdf.open(path), 1):
        lines.append(f"[p. {number}]")
        for line in page.get_text().splitlines():
            line = line.strip()
            # Drop running page numbers and empty lines.
            if line and not line.isdigit():
                lines.append(line)
    chosen = [line for start, end in THESIS_SECTIONS for line in section(lines, start, end)]
    # Re-join PDF line wraps into paragraphs so the text is denser for the model.
    text = "\n".join(chosen)
    text = re.sub(r"(?<![.:;!?»\]])\n(?!\[p\. )", " ", text)
    return re.sub(r"[ \t]+", " ", text)


def main():
    if len(sys.argv) < 4:
        sys.exit(__doc__)
    slides_path, qa_path, thesis_path = sys.argv[1:4]
    out = Path(sys.argv[4] if len(sys.argv) > 4 else "knowledge")
    out.mkdir(parents=True, exist_ok=True)

    slides_md, note_pairs = extract_slides(slides_path)
    doc_pairs = extract_qa_doc(qa_path)
    thesis_md = extract_thesis(thesis_path)

    (out / "slides.md").write_text(slides_md)
    (out / "qa.json").write_text(json.dumps(doc_pairs + note_pairs, ensure_ascii=False, indent=1))
    (out / "thesis.md").write_text(thesis_md)
    print(
        f"slides.md: {len(slides_md):,} chars; qa.json: {len(doc_pairs)} from the Q&A document + "
        f"{len(note_pairs)} from slide notes; thesis.md: {len(thesis_md):,} chars -> {out}"
    )


if __name__ == "__main__":
    main()
